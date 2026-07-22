"""
One-time builder: reads the source matrix workbook and writes the two
machine-readable reference assets injected into the judge's system prompt.

Source: D:\\tritorc\\data_enrich\\Designation _ Problems _ Obstacles.xlsx
Outputs (this folder): targeting.json, icp_companies.json

Re-run whenever the source spreadsheet's "Targeting" or "ICP List" sheets change:
    python build_reference_data.py
"""

import json
from pathlib import Path

import openpyxl

SOURCE = Path(r"D:\tritorc\data_enrich\Designation _ Problems _ Obstacles.xlsx")
OUT_DIR = Path(__file__).parent


def build_targeting(wb):
    ws = wb["Targeting"]
    rows = list(ws.iter_rows(min_row=1, max_row=57, values_only=True))
    header = rows[0]
    function_groups = [h.strip() if h else None for h in header[1:8]]

    matrix = {}
    current_type = None
    for row in rows[1:]:
        customer_type = row[0]
        if customer_type:
            current_type = customer_type.strip()
            matrix[current_type] = {fg: [] for fg in function_groups if fg}
        if current_type is None:
            continue
        for fg, value in zip(function_groups, row[1:8]):
            if fg and value:
                title = str(value).replace("�", "-").strip()
                if title and title not in matrix[current_type][fg]:
                    matrix[current_type][fg].append(title)

    return {
        ctype: {fg: titles for fg, titles in groups.items() if titles}
        for ctype, groups in matrix.items()
    }


def build_icp(wb):
    ws = wb["ICP List"]
    rows = list(ws.iter_rows(min_row=1, max_row=33, values_only=True))
    industries_row = rows[1]
    industries = [(i, name.strip()) for i, name in enumerate(industries_row) if name]

    icp = {}
    use_cases = {}
    for col_idx, industry in industries:
        column = [row[col_idx] for row in rows[2:]]
        blocks = []
        current_label = None
        current_items = []
        for raw in column + [None]:
            if raw is None:
                if current_label is not None:
                    blocks.append((current_label, current_items))
                current_label = None
                current_items = []
                continue
            text = str(raw).replace("�", "-").strip()
            if not text:
                continue
            if current_label is None:
                current_label = text
                current_items = []
            else:
                current_items.append(text)

        categories = {}
        industry_use_cases = []
        for label, items in blocks:
            if label in ("Use case for Tritorc", "Use Case"):
                industry_use_cases.extend(items)
                continue
            if items:
                categories[label] = items
        icp[industry] = categories
        if industry_use_cases:
            use_cases[industry] = industry_use_cases

    return icp, use_cases


def main():
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    targeting = build_targeting(wb)

    (OUT_DIR / "targeting.json").write_text(
        json.dumps(targeting, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print("targeting.json:", len(targeting), "customer types")

    if "ICP List" not in wb.sheetnames:
        print("no 'ICP List' sheet in the source workbook — leaving icp_companies.json untouched")
        return

    icp, use_cases = build_icp(wb)
    icp_out = {"industries": icp}
    if use_cases:
        icp_out["bolting_use_cases"] = use_cases
    (OUT_DIR / "icp_companies.json").write_text(
        json.dumps(icp_out, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print("icp_companies.json:", len(icp), "industries")


if __name__ == "__main__":
    main()
