"""Schema-locked Excel writer. Column order is defined once here as
(header, key) tuples — a vendor field rename is a one-line change."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

CONTACT_SCHEMA: list[tuple[str, str]] = [
    ("First Name", "first_name"),
    ("Last Name", "last_name"),
    ("Full Name", "full_name"),
    ("Title", "title"),
    ("Seniority", "seniority"),
    ("Person Linkedin Url", "linkedin_url"),
    ("City", "city"),
    ("State", "state"),
    ("Country", "country"),
    ("Email", "email"),
    ("Email Status", "email_status"),
    ("Company Name", "company_name"),
    ("Company Domain", "company_domain"),
    ("Website", "website"),
    ("Industry", "industry"),
    ("# Employees", "employees"),
    ("Annual Revenue", "annual_revenue"),
    ("Total Funding", "total_funding"),
    ("Company Phone", "company_phone"),
    ("Company Linkedin Url", "company_linkedin_url"),
    ("Company Street", "company_street"),
    ("Company City", "company_city"),
    ("Company Postal Code", "company_postal_code"),
    ("Company State", "company_state"),
    ("Company Country", "company_country"),
    ("Company Founded Year", "company_founded_year"),
    ("_Customer Type", "customer_type"),
    ("_Size Tier", "size_tier"),
    ("_Verdict", "verdict"),
    ("_Confidence", "confidence"),
    ("_Matched Designation", "matched_designation"),
    ("_Function Group", "function_group"),
    ("_Reasoning", "reasoning"),
    ("_Prompt Version", "prompt_version"),
    ("_Overridden", "overridden"),
    ("_Revealed", "revealed"),
    ("_Source", "source"),
]

RUN_LOG_SCHEMA: list[tuple[str, str]] = [
    ("Company Input Name", "company_input_name"),
    ("Domain", "domain"),
    ("Resolved", "resolved"),
    ("Resolution Source", "resolution_source"),
    ("Customer Type", "customer_type"),
    ("Size Tier", "size_tier"),
    ("Employees Found", "employees_found"),
    ("Judged", "judged"),
    ("Approved", "approved"),
    ("Revealed", "revealed"),
    ("Credits Spent", "credits_spent"),
    ("Status", "status"),
    ("Notes", "notes"),
]

NEEDS_REVIEW_SCHEMA: list[tuple[str, str]] = [
    ("Company Input Name", "company_input_name"),
    ("Domain", "domain"),
    ("Reason", "reason"),
]

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill(start_color="1F4E5F", end_color="1F4E5F", fill_type="solid")
_BODY_FONT = Font(name="Calibri", size=11)


def _write_sheet(ws, schema: list[tuple[str, str]], rows: list[dict]) -> None:
    ws.append([header for header, _ in schema])
    for col_idx in range(1, len(schema) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        ws.append([row.get(key) if row.get(key) is not None else "" for _, key in schema])

    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            cell.font = _BODY_FONT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col_idx, (header, _) in enumerate(schema, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = max(12, min(40, len(header) + 4))


def write_output_excel(contacts: list[dict], run_log: list[dict],
                       needs_review: list[dict], output_path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Contacts"
    _write_sheet(ws, CONTACT_SCHEMA, contacts)
    _write_sheet(wb.create_sheet("Run Log"), RUN_LOG_SCHEMA, run_log)
    if needs_review:
        _write_sheet(wb.create_sheet("Needs Review"), NEEDS_REVIEW_SCHEMA, needs_review)
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
