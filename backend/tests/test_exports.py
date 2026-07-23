"""search_only export: every swept person, any verdict, no credits spent."""

import openpyxl
import pytest
from mongomock_motor import AsyncMongoMockClient
from beanie import init_beanie

from app.api.exports import export_run
from app.db.session import ALL_DOCUMENTS
from app.models.company import Company
from app.models.employee import Employee, Verdict, Reveal
from app.models.run import Run


@pytest.fixture(autouse=True)
async def db():
    client = AsyncMongoMockClient()
    await init_beanie(database=client["test"], document_models=ALL_DOCUMENTS)
    yield


async def _seed():
    run = Run(filename="t.xlsx", companies_total=1)
    await run.insert()
    company = Company(run_id=run.id, input_name="Acme", website="acme.com", domain="acme.com")
    await company.insert()

    enriched = Employee(run_id=run.id, company_id=company.id, apollo_id="p1",
                        full_name="Jane Doe", title="Reliability Engineer",
                        verdict=Verdict(verdict="enrich", confidence=0.95, reasoning="fixture"),
                        reveal=Reveal(status="revealed", email="jane@acme.com", engine="apollo", credits=1))
    await enriched.insert()
    skipped = Employee(run_id=run.id, company_id=company.id, apollo_id="p2",
                       full_name="Bob Smith", title="Marketing Intern",
                       verdict=Verdict(verdict="skip", confidence=0.9, reasoning="fixture"))
    await skipped.insert()
    unjudged = Employee(run_id=run.id, company_id=company.id, apollo_id="p3",
                        full_name="Ann Lee", title="Warehouse Associate")
    await unjudged.insert()
    return run


async def test_search_only_includes_everyone_no_email():
    run = await _seed()
    response = await export_run(run.id, search_only=True)

    wb = openpyxl.load_workbook(response.path)
    ws = wb["Contacts"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    names = {row[2] for row in rows}  # Full Name column
    assert names == {"Jane Doe", "Bob Smith", "Ann Lee"}

    header = [c.value for c in ws[1]]
    name_col = header.index("Full Name")
    email_col = header.index("Email")
    source_col = header.index("_Source")
    by_name = {row[name_col]: row for row in rows}

    # Never-revealed people: no email, clearly labeled search-only.
    for name in ("Bob Smith", "Ann Lee"):
        assert not by_name[name][email_col]  # openpyxl reads a written "" back as None
        assert by_name[name][source_col] == "search_only (not revealed)"
    # Already-revealed people keep their real data (search_only widens
    # inclusion, it doesn't redact people who were legitimately revealed).
    assert by_name["Jane Doe"][email_col] == "jane@acme.com"


async def test_search_only_does_not_touch_export_path():
    run = await _seed()
    await export_run(run.id, search_only=True)
    run = await Run.get(run.id)
    assert run.export_path is None


async def test_default_export_excludes_unapproved():
    run = await _seed()
    response = await export_run(run.id)
    wb = openpyxl.load_workbook(response.path)
    ws = wb["Contacts"]
    names = {row[2] for row in ws.iter_rows(min_row=2, values_only=True)}
    assert names == {"Jane Doe"}


async def test_search_only_keeps_every_person_when_names_are_unrevealed():
    # Regression: Apollo's free preview never returns a real full_name (only
    # first_name + an obfuscated last name), so several approved-but-unrevealed
    # people at the same company used to collapse into a single dedup row.
    run = Run(filename="t.xlsx", companies_total=1)
    await run.insert()
    company = Company(run_id=run.id, input_name="Acme", website="acme.com", domain="acme.com")
    await company.insert()

    for i in range(4):
        emp = Employee(run_id=run.id, company_id=company.id, apollo_id=f"p{i}",
                       first_name=f"Person{i}", last_name="O***s",
                       full_name=None, title="Welder",
                       verdict=Verdict(verdict="enrich", confidence=0.9, reasoning="fixture"))
        await emp.insert()

    response = await export_run(run.id, search_only=True)
    wb = openpyxl.load_workbook(response.path)
    ws = wb["Contacts"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    header = [c.value for c in ws[1]]
    first_name_col = header.index("First Name")
    assert len(rows) == 4
    assert {row[first_name_col] for row in rows} == {"Person0", "Person1", "Person2", "Person3"}
